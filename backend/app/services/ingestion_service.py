import os
import re
import json
import logging
from typing import List, Dict, Any
from sqlalchemy.orm import Session

from app.models.document import Document, DocumentChunk
from app.models.knowledge import Topic, Concept, TopicRelationship
from app.models.quiz import Question
from app.models.flashcard import Flashcard
from app.models.pyq import PYQQuestion
from app.models.mastery import MasteryScore
from app.ai.llm_provider import llm
from app.core.json_utils import parse_json_safely

logger = logging.getLogger(__name__)

class IngestionService:
    def clean_educational_text(self, text: str) -> str:
        if not text:
            return ""
        pdf_obj_pattern = re.compile(r'\d+\s+\d+\s+obj', re.IGNORECASE)
        pdf_tag_pattern = re.compile(r'/(MediaBox|Linearized|FlateDecode|Font|Catalog|Pages|ProcSet|ExtGState|Type|Filter|Length|Parent|Resources|Encoding|Widths|FontDescriptor|Root|Info)\b', re.IGNORECASE)
        
        lines = text.split("\n")
        cleaned_lines = []
        font_stopwords = {
            "times new roman", "ms mincho", "arial", "calibri", "courier new", "tahoma",
            "verdana", "georgia", "helvetica", "symbol", "wingdings", "rectangle",
            "default design", "header placeholder", "date placeholder", "notes placeholder",
            "footer placeholder", "slide number placeholder", "title placeholder", "body placeholder"
        }
        
        for line in lines:
            l_trim = line.strip()
            if not l_trim:
                continue
            if (pdf_obj_pattern.search(l_trim) or 
                pdf_tag_pattern.search(l_trim) or 
                l_trim.startswith("%PDF") or 
                l_trim.startswith("endobj") or 
                l_trim.startswith("stream") or 
                l_trim.startswith("endstream") or
                l_trim.startswith("RSoP") or
                l_trim.startswith("Root Entry") or
                l_trim.startswith("<<") or
                l_trim.endswith(">>")):
                continue
            l_lower = l_trim.lower()
            if any(sw in l_lower for sw in font_stopwords) or l_lower.startswith("rectangle ") or l_lower.startswith("oval "):
                continue
            cleaned_lines.append(l_trim)
        return "\n".join(cleaned_lines)

    def extract_text(self, file_path: str, file_type: str) -> List[Dict[str, Any]]:
        pages = []
        if not file_path or not os.path.exists(file_path):
            return []

        ext = file_type.lower()

        if ext == "pdf":
            try:
                import pypdf
                reader = pypdf.PdfReader(file_path)
                for idx, page in enumerate(reader.pages):
                    text = page.extract_text() or ""
                    cleaned = self.clean_educational_text(text)
                    if cleaned.strip():
                        pages.append({"page_number": idx + 1, "text": cleaned.strip()})
            except Exception as e:
                logger.error(f"Error extracting PDF text: {e}")

        elif ext == "docx":
            try:
                import docx
                doc = docx.Document(file_path)
                text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
                cleaned = self.clean_educational_text(text)
                if cleaned.strip():
                    pages.append({"page_number": 1, "text": cleaned.strip()})
            except Exception as e:
                logger.error(f"Error extracting DOCX text: {e}")

        elif ext in ["pptx", "ppt"]:
            try:
                import pptx
                prs = pptx.Presentation(file_path)
                for idx, slide in enumerate(prs.slides):
                    slide_texts = []
                    for shape in slide.shapes:
                        if hasattr(shape, "text") and shape.text.strip():
                            slide_texts.append(shape.text.strip())
                        elif shape.has_table:
                            for row in shape.table.rows:
                                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                                if row_text:
                                    slide_texts.append(row_text)
                    full_slide_text = "\n".join(slide_texts)
                    cleaned = self.clean_educational_text(full_slide_text)
                    if cleaned.strip():
                        pages.append({"page_number": idx + 1, "text": cleaned.strip()})
            except Exception as e:
                logger.error(f"Error extracting PPTX: {e}")

        # Fallback binary / plain text reader for unhandled formats or binary legacy PPT files
        if not pages:
            try:
                with open(file_path, "rb") as f:
                    raw_data = f.read()
                    extracted_lines = []

                    try:
                        utf16_decoded = raw_data.decode("utf-16le", errors="ignore")
                        utf16_matches = re.findall(r'[\x20-\x7E\s]{4,}', utf16_decoded)
                        for m in utf16_matches:
                            cleaned = m.strip()
                            if len(cleaned) > 3 and len(cleaned.split()) >= 2:
                                extracted_lines.append(cleaned)
                    except Exception:
                        pass

                    ascii_matches = re.findall(rb'[\x20-\x7E\s]{4,}', raw_data)
                    for m in ascii_matches:
                        cleaned = m.decode("ascii", errors="ignore").strip()
                        if len(cleaned) > 3 and len(cleaned.split()) >= 2:
                            if cleaned not in extracted_lines:
                                extracted_lines.append(cleaned)

                    cleaned_full = self.clean_educational_text("\n".join(extracted_lines[:500]))
                    if cleaned_full.strip():
                        pages.append({"page_number": 1, "text": cleaned_full.strip()})
            except Exception as e:
                logger.error(f"Fallback binary text reader error: {e}")

        return pages

    def chunk_text(self, pages: List[Dict[str, Any]], chunk_size: int = 400, overlap: int = 50) -> List[Dict[str, Any]]:
        chunks = []
        chunk_idx = 0
        
        for p in pages:
            page_num = p.get("page_number", 1)
            text = p.get("text", "")
            words = text.split()
            
            if not words:
                continue

            for i in range(0, len(words), max(1, chunk_size - overlap)):
                chunk_words = words[i:i + chunk_size]
                chunk_text = " ".join(chunk_words)
                
                heading = f"Page {page_num}: Section {chunk_idx + 1}"
                lines = [l.strip() for l in chunk_text.split("\n") if l.strip()]
                if lines and len(lines[0]) < 80 and not lines[0].startswith("http"):
                    heading = lines[0]

                chunks.append({
                    "chunk_index": chunk_idx,
                    "page_number": page_num,
                    "section_heading": heading,
                    "content": chunk_text
                })
                chunk_idx += 1

        return chunks

    def process_document(self, db: Session, document_id: int):
        doc = db.query(Document).filter(Document.id == document_id).first()
        if not doc:
            return

        try:
            doc.status = "extracting"
            db.commit()

            # 1. Full PDF / Document Text Extraction
            pages = self.extract_text(doc.file_path or "", doc.file_type)
            if not pages:
                doc.status = "failed"
                doc.error_message = "No readable text content found in uploaded document."
                db.commit()
                return

            doc.status = "structure_detected"
            db.commit()

            chunks_data = self.chunk_text(pages)
            doc.status = "embedding"
            db.commit()

            # Save chunks to DB
            for c in chunks_data:
                chunk_obj = DocumentChunk(
                    document_id=doc.id,
                    chunk_index=c["chunk_index"],
                    page_number=c["page_number"],
                    section_heading=c["section_heading"],
                    content=c["content"],
                    embedding_ref=f"chunk_{doc.id}_{c['chunk_index']}"
                )
                db.add(chunk_obj)
            db.commit()

            # Combine full PDF text (up to 120,000 characters) to pass to Groq API for full document content analysis
            full_pdf_text = "\n\n".join([f"[Page {p['page_number']}]\n{p['text']}" for p in pages])
            truncated_text = full_pdf_text[:120000]

            clean_doc_title = doc.title.split(".")[0].replace("_", " ").replace("-", " ").title()

            # 2. Call Groq API for Comprehensive PDF Content Analysis, Grounded Questions & PYQ Analysis
            system_prompt = (
                "You are an expert educational AI system. Analyze the provided study material PDF text and generate structured learning content grounded strictly in the PDF text.\n"
                "DO NOT generate random or placeholder questions. All questions, options, definitions, and PYQ analysis MUST be directly traceable to the facts, formulas, and topics in the uploaded PDF text.\n\n"
                "Output JSON format:\n"
                "{\n"
                '  "pyq_analysis": {\n'
                '    "frequently_asked_topics": ["Topic 1", "Topic 2"],\n'
                '    "repeated_concepts": ["Concept A", "Concept B"],\n'
                '    "question_patterns": ["Derivation questions", "Numerical applications"],\n'
                '    "difficulty_patterns": "Summary of difficulty distribution across exam questions"\n'
                '  },\n'
                '  "modules": [\n'
                '    {\n'
                '      "name": "Module Name / Topic",\n'
                '      "description": "Brief summary of concepts in this module",\n'
                '      "mcqs": [\n'
                '        {\n'
                '          "prompt": "Clear MCQ question prompt based on PDF text",\n'
                '          "options": ["Correct choice", "Distractor 1", "Distractor 2", "Distractor 3"],\n'
                '          "correct_answer": "Correct choice",\n'
                '          "explanation": "Explanation grounded in PDF text",\n'
                '          "difficulty": "easy"|"medium"|"hard"\n'
                '        }\n'
                '      ],\n'
                '      "short_answers": [\n'
                '        {\n'
                '          "prompt": "Short answer question prompt",\n'
                '          "correct_answer": "Exact reference answer based on PDF",\n'
                '          "explanation": "Explanation",\n'
                '          "difficulty": "medium"\n'
                '        }\n'
                '      ],\n'
                '      "long_answers": [\n'
                '        {\n'
                '          "prompt": "Long analytical answer question prompt",\n'
                '          "correct_answer": "Comprehensive answer based on PDF text",\n'
                '          "explanation": "Detailed explanation",\n'
                '          "difficulty": "hard"\n'
                '        }\n'
                '      ],\n'
                '      "flashcards": [\n'
                '        {\n'
                '          "card_type": "concept"|"formula"|"definition",\n'
                '          "front": "Term or concept front",\n'
                '          "back": "Detailed definition or formula from PDF"\n'
                '        }\n'
                '      ],\n'
                '      "pyqs": [\n'
                '        {\n'
                '          "year": 2024,\n'
                '          "exam_name": "Exam Analysis",\n'
                '          "question_text": "Grounded practice or historical exam question",\n'
                '          "marks": 10,\n'
                '          "difficulty": "medium"|"hard"\n'
                '        }\n'
                '      ]\n'
                '    }\n'
                '  ]\n'
                '}'
            )

            user_prompt = f"Document Title: {doc.title}\n\nFull Uploaded PDF Text Content:\n{truncated_text}"

            raw_analysis = llm.generate_completion(system_prompt, user_prompt, json_mode=True)
            
            parsed = parse_json_safely(raw_analysis, fallback_default={})
            modules_list = parsed.get("modules", [])
            pyq_analysis_data = parsed.get("pyq_analysis", {})

            # Save pyq_analysis into document metadata
            meta = doc.doc_metadata or {}
            if isinstance(meta, dict):
                meta["pyq_analysis"] = pyq_analysis_data
                doc.doc_metadata = meta
                db.commit()

            created_topics: List[Topic] = []

            if modules_list:
                for idx, mod in enumerate(modules_list[:8]):
                    mod_name = mod.get("name", f"Module {idx + 1}: {clean_doc_title}")
                    mod_desc = mod.get("description", f"Grounded module from {doc.title}")

                    topic = db.query(Topic).filter(Topic.course_id == doc.course_id, Topic.name == mod_name).first()
                    if not topic:
                        topic = Topic(
                            course_id=doc.course_id,
                            name=mod_name,
                            description=mod_desc
                        )
                        db.add(topic)
                        db.commit()
                        db.refresh(topic)

                    created_topics.append(topic)

                    # Initial Mastery Score
                    mastery = db.query(MasteryScore).filter(
                        MasteryScore.course_id == doc.course_id,
                        MasteryScore.topic_id == topic.id
                    ).first()
                    if not mastery:
                        db.add(MasteryScore(
                            user_id=1,
                            course_id=doc.course_id,
                            topic_id=topic.id,
                            score=60.0 + (idx * 4.0) % 30.0,
                            confidence=80.0
                        ))

                    # Insert Grounded MCQs
                    for q in mod.get("mcqs", []):
                        if q.get("prompt"):
                            db.add(Question(
                                topic_id=topic.id,
                                question_type="mcq",
                                difficulty=q.get("difficulty", "medium"),
                                prompt=q.get("prompt"),
                                options=q.get("options", []),
                                correct_answer=q.get("correct_answer", ""),
                                explanation=q.get("explanation", f"Grounded in '{doc.title}'.")
                            ))

                    # Insert Grounded Short Answer Questions
                    for q in mod.get("short_answers", []):
                        if q.get("prompt"):
                            db.add(Question(
                                topic_id=topic.id,
                                question_type="short_answer",
                                difficulty=q.get("difficulty", "medium"),
                                prompt=q.get("prompt"),
                                options=None,
                                correct_answer=q.get("correct_answer", ""),
                                explanation=q.get("explanation", f"Grounded in '{doc.title}'.")
                            ))

                    # Insert Grounded Long Answer Questions
                    for q in mod.get("long_answers", []):
                        if q.get("prompt"):
                            db.add(Question(
                                topic_id=topic.id,
                                question_type="long_answer",
                                difficulty=q.get("difficulty", "hard"),
                                prompt=q.get("prompt"),
                                options=None,
                                correct_answer=q.get("correct_answer", ""),
                                explanation=q.get("explanation", f"Grounded in '{doc.title}'.")
                            ))

                    # Insert Grounded Flashcards
                    for fc in mod.get("flashcards", []):
                        if fc.get("front") and fc.get("back"):
                            db.add(Flashcard(
                                course_id=doc.course_id,
                                topic_id=topic.id,
                                card_type=fc.get("card_type", "concept"),
                                front=fc.get("front"),
                                back=fc.get("back")
                            ))

                    # Insert Grounded PYQs
                    for pyq_item in mod.get("pyqs", []):
                        if pyq_item.get("question_text"):
                            db.add(PYQQuestion(
                                course_id=doc.course_id,
                                topic_id=topic.id,
                                year=pyq_item.get("year", 2024),
                                exam_name=pyq_item.get("exam_name", "Exam Analysis"),
                                question_text=pyq_item.get("question_text"),
                                marks=pyq_item.get("marks", 10),
                                difficulty=pyq_item.get("difficulty", "medium"),
                                question_type="long_answer",
                                occurrences_count=3
                            ))

            # Build Knowledge Graph Edges between consecutive modules
            for i in range(len(created_topics) - 1):
                t1 = created_topics[i]
                t2 = created_topics[i + 1]
                rel = db.query(TopicRelationship).filter(
                    TopicRelationship.source_topic_id == t1.id,
                    TopicRelationship.target_topic_id == t2.id
                ).first()
                if not rel:
                    db.add(TopicRelationship(
                        source_topic_id=t1.id,
                        target_topic_id=t2.id,
                        relationship_type="prerequisite"
                    ))

            doc.status = "ready"
            db.commit()

        except Exception as e:
            logger.error(f"Failed to process document {document_id}: {e}", exc_info=True)
            doc.status = "failed"
            doc.error_message = str(e)
            db.commit()

ingestion_service = IngestionService()


